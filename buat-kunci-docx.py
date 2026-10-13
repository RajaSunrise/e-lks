#!/usr/bin/env python3
"""Membuat DOCX Kunci Jawaban E-LKS 1-4 + Kuis (guru/peneliti). Jalankan: python3 buat-kunci-docx.py"""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HIJAU = "0E2F21"
ABU = "F2F9F3"

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin, s.bottom_margin = Cm(1.5), Cm(2.0)

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(10)


def judul(teks):
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(18), RGBColor(0x16, 0x65, 0x34)
    return p


def h1(teks):
    p = doc.add_paragraph()
    r = p.add_run(teks)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), RGBColor(0x16, 0x65, 0x34)
    return p


def teks(body, merah=False, kecil=False):
    p = doc.add_paragraph()
    r = p.add_run(body)
    if merah:
        r.bold, r.font.color.rgb = True, RGBColor(0xDC, 0x26, 0x26)
    if kecil:
        r.font.size = Pt(9)
    return p


def arsir(sel, warna):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), warna)
    sel._tc.get_or_add_tcPr().append(shd)


def tabel(baris):
    t = doc.add_table(rows=len(baris), cols=len(baris[0]))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(baris):
        for j, isi in enumerate(row):
            sel = t.cell(i, j)
            sel.text = ""
            p = sel.paragraphs[0]
            r = p.add_run(isi)
            r.font.size = Pt(9)
            if i == 0:
                r.bold, r.font.color.rgb = True, RGBColor(0xFF, 0xFF, 0xFF)
                arsir(sel, HIJAU)
            elif i % 2 == 0:
                arsir(sel, ABU)
    doc.add_paragraph()
    return t


def footer():
    f = doc.sections[0].footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run("Kunci E-LKS Kehati PBL - MAN Lombok Barat - RAHASIA GURU - Hal. ")
    r.font.size = Pt(8)
    fld1, instr, fld2 = OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld2.set(qn("w:fldCharType"), "end")
    f._p.append(fld1)
    f._p.append(instr)
    f._p.append(fld2)


judul("KUNCI JAWABAN & RUBRIK PENILAIAN")
teks("E-LKS 1-4 Keanekaragaman Hayati Berbasis PBL")
teks("Biologi Kelas X - MAN Lombok Barat - Berpikir Kritis")
teks("RAHASIA GURU/PENELITI - jangan disebar ke siswa", merah=True)

tabel([
    ["Komponen", "Bobot", "Keterangan"],
    ["E-LKS 1 (tingkat, persebaran, manfaat)", "10", "Q1=2, Q2=4, Q3=2, Q4=2"],
    ["E-LKS 2 (faktor hilang)", "10", "Kasus+bukti=4, dampak=2, tunggal/asumsi=2, tindak=2"],
    ["E-LKS 3 (pelestarian)", "10", "Ancaman=2, banding=3, solusi=2, alasan=1, peran=2"],
    ["E-LKS 4 (klasifikasi)", "10", "Ciri=2, dikotomi=3, kingdom=3, takson=2"],
    ["Kuis PG (10 soal)", "40", "4 poin per soal"],
    ["Isian singkat (2 soal)", "10", "5 poin per soal"],
    ["Uraian (2 soal)", "10", "5 poin per soal"],
    ["TOTAL", "100", "KKM 75"],
])

h1("E-LKS 1 - Tingkat, Persebaran, Manfaat (Pertemuan 2)")
tabel([
    ["No", "Pertanyaan", "Kunci (bobot)", "Pembahasan"],
    ["1", "Inti permasalahan wacana Lombok-Bali?", "B (2)", "Inti: mengapa fauna Lombok-Bali berbeda + penyebab serta manfaat kehati."],
    ["2", "Contoh gen, jenis, ekosistem + dasar beda", "Rubrik (4)", "Kata gen+jenis+ekosistem (2); contoh benar mis. padi/komodo/hutan (1); jawaban >60 huruf (1)."],
    ["3", "Sumber paling dipercaya?", "C (2)", "Jurnal/buku teks/peneliti - berbasis data dan telaah ahli."],
    ["4", "Simpulkan penyebab beda persebaran", "Rubrik (2)", "Garis Wallace/selat/Lombok-Bali (1); iklim/geografi/tanah/adaptasi (1)."],
])

h1("E-LKS 2 - Faktor Hilangnya Kehati (Pertemuan 3)")
tabel([
    ["No", "Pertanyaan", "Kunci (bobot)", "Pembahasan"],
    ["1", "Pilih 1 dari 6 kasus + bukti", "Pilih=1, bukti=3", "Bukti dinilai dari kata: bukti/data/amatan/lihat/foto/ukur/catat (maks 3); min. 10 huruf."],
    ["2", "Dampak bagi ekosistem dan manusia", "Rubrik (2)", "Dampak ekosistem: rantai/punah/habitat (1); manusia: ekonomi/nelayan/kesehatan (1)."],
    ["3", "Penyebab tunggal? + asumsi", "B+asumsi (2)", "B = saling terkait. Asumsi >15 huruf (1)."],
    ["4", "Tindakan pencegahan", "Rubrik (2)", "Kata: cegah/tanam/bersih/hukum/edukasi/daur (maks 2); min. 10 huruf."],
])
teks("6 kasus: (1) alih fungsi mangrove, (2) sampah plastik & pencemaran, (3) pemutihan karang, (4) perburuan berlebihan, (5) spesies pendatang, (6) monokultur & penebangan.", kecil=True)

h1("E-LKS 3 - Upaya Pelestarian (Pertemuan 4)")
tabel([
    ["No", "Pertanyaan", "Kunci (bobot)", "Pembahasan"],
    ["1", "Ancaman utama", "Rubrik (2)", "Min. 2 dari: alih fungsi, buru, sampah, cemar, iklim, invasif, tebang, monokultur, limbah, kebakaran."],
    ["2", "Bandingkan in-situ vs ex-situ", "Rubrik (3)", "Menyebut in-situ (1); ex-situ (1); kelebihan/keterbatasan/contoh (1)."],
    ["3", "Solusi + alasan", "C+alasan (3)", "C = kombinasi keduanya (2); alasan >20 huruf (1)."],
    ["4", "Peran yang dapat dilakukan", "Rubrik (2)", "Kata: tanam/bersih/edukasi/hemat/lapor/pilah (maks 2); min. 10 huruf."],
])

h1("E-LKS 4 - Klasifikasi (Pertemuan 5)")
tabel([
    ["No", "Pertanyaan", "Kunci (bobot)", "Pembahasan"],
    ["1", "Ciri dasar pengelompokan?", "B (2)", "Struktur sel, dinding sel, dan cara memperoleh makanan."],
    ["2", "Susun kunci dikotomi", "Rubrik (3)", "Format pasangan 1a/1b (1); min. 2 ciri pembeda (1); >40 huruf (1)."],
    ["3", "Kingdom 6 organisme", "Tabel (3)", "0,5 per benar. Lihat tabel kingdom di bawah."],
    ["4", "Definisi takson", "Rubrik (2)", "Takson/kelompok/tingkat (1); contoh tingkatan mis. kingdom-filum-kelas (1)."],
])
tabel([
    ["Organisme", "Kingdom", "Alasan singkat"],
    ["Bakteri", "Monera", "Prokariot, uniseluler"],
    ["Jamur", "Fungi", "Dinding kitin, pengurai"],
    ["Lumut", "Plantae", "Bersel banyak, berklorofil, berspora"],
    ["Paku", "Plantae", "Berpembuluh, berspora"],
    ["Mawar", "Plantae", "Berbiji tertutup, berbunga"],
    ["Kucing", "Animalia", "Tanpa dinding sel, bergerak aktif"],
])

h1("Kuis Pilihan Ganda (10 x 4 = 40)")
tabel([
    ["No", "Kunci", "Pembahasan singkat"],
    ["1", "B", "Mawar beda warna satu spesies = tingkat gen."],
    ["2", "C", "Harimau vs singa = spesies berbeda (jenis)."],
    ["3", "B", "Hutan vs karang = beda ekosistem."],
    ["4", "C", "Komodo-anoa-babirusa = Wallacea/peralihan."],
    ["5", "B", "TN Komodo/Rinjani = habitat asli = in-situ."],
    ["6", "C", "Alih fungsi habitat/deforestasi = penyebab utama."],
    ["7", "B", "Rafflesia: endemik, langka, terancam punah."],
    ["8", "B", "IR8 vs Ciherang = satu spesies beda varietas."],
    ["9", "C", "Sawah-waduk-kebun = ekosistem buatan."],
    ["10", "B", "Jamur: dinding kitin + pengurai = Fungi."],
])

h1("Isian Singkat (2 x 5 = 10)")
tabel([
    ["No", "Kunci", "Bobot"],
    ["11", "Tingkat gen (5; jawaban relevan lain = 2)", "5"],
    ["12", "Wallacea / peralihan / tengah (5; relevan = 2)", "5"],
])

h1("Uraian (2 x 5 = 10)")
tabel([
    ["No", "Rubrik", "Bobot"],
    ["13", "In-situ/habitat asli (1); ex-situ (1); contoh TN/kebun binatang/bank benih (1); >40 huruf (1); >100 huruf (1).", "5"],
    ["14", "Ancaman maks 3 (1 tiap kata); solusi maks 2.", "5"],
])

h1("Pemetaan Aspek Berpikir Kritis")
tabel([
    ["Aspek", "Sumber butir", "Maks"],
    ["Interpretasi", "E-LKS 1 (Q1+Q2) + E-LKS 4 takson + Isian", "18"],
    ["Analisis", "E-LKS 2 + E-LKS 4 + PG 1,2,3,8,9,10", "38"],
    ["Evaluasi", "E-LKS 1 Q3 + E-LKS 2 Q3 + E-LKS 3 + PG 4,5,6,7", "27"],
    ["Inferensi", "E-LKS 1 Q4 + E-LKS 2 tindak + E-LKS 3 alasan/peran + Uraian 14", "12"],
    ["Eksplanasi", "Uraian 13", "5"],
])
teks("Kategori: 85+ Sangat Kritis; 75-84 Kritis; 60-74 Cukup Kritis; di bawah 60 Kurang Kritis. Regulasi diri dinilai lewat Refleksi (tidak diskor).", kecil=True)

footer()
NAMA = "Kunci-Jawaban-E-LKS-Kehati.docx"
doc.save(NAMA)
print("OK:", NAMA)
